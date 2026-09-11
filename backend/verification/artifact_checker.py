"""Artifact checking (Phase 10 + Phase 11).

The review was specific about this one: a deck is not verified because the
generator exited 0. It is verified when something opens the file and finds
slides in it.

So there are two layers here, and the split is deliberate.

**Host layer - integrity and container structure.** Uses only the standard
library: the file exists, its size matches what was recorded, its SHA-256
still matches the digest taken as the bytes left the sandbox, and the
container is well formed (OOXML files are ZIPs with required parts; a PDF
starts with ``%PDF-`` and ends with ``%%EOF``). A mismatch here is the
serious case - it means the artifact was truncated or altered after
generation - so it blocks completion.

**Sandbox layer - open it with the real library.** ``python-pptx`` /
``python-docx`` / ``openpyxl`` are not host dependencies, and installing them
on the host to inspect model-driven output would put a parser for untrusted
files in the API process. So the artifact is base64-ed into the sandbox and
opened there, by a constant reviewed script, in the same locked-down container
that produced it. That is where "no empty slides", "sheets contain rows" and
the overflow estimate come from.

When the sandbox is not configured, or the artifact is too large to ship in as
an input, the deep layer is reported as a **WARNING** rather than skipped
quietly. An artifact whose contents nobody opened must not read as verified.

PDF deep inspection is honestly absent: the documents image carries a PDF
*writer* (ReportLab), not a reader. Page count is derived structurally on the
host and reported as such.
"""

from __future__ import annotations

import base64
import hashlib
import json
import re
import time
import zipfile
from pathlib import Path
from typing import Any

from backend.verification.base import CheckResult, CheckStatus, VerificationInput

#: Above this the base64 payload stops being a reasonable sandbox input.
_MAX_SANDBOX_INSPECT_BYTES = 6 * 1024 * 1024

#: One required part per OOXML type - absence means the container is not what
#: it claims to be, whatever the extension says.
_OOXML_PARTS = {
    "pptx": "ppt/presentation.xml",
    "docx": "word/document.xml",
    "xlsx": "xl/workbook.xml",
}
_OOXML_TYPES = tuple(_OOXML_PARTS)

_PDF_PAGE = re.compile(rb"/Type\s*/Page[^s]")

#: Runs inside the sandbox. A constant, reviewed script - not model output -
#: so it is data as far as the sandbox policy is concerned.
_INSPECT_SCRIPT = r'''
import base64
import json
import sys

kind = sys.argv[1]
ext = {"pptx": ".pptx", "docx": ".docx", "xlsx": ".xlsx"}[kind]
path = "/workspace/tmp/artifact" + ext

with open("inputs/artifact.b64", "r", encoding="utf-8") as handle:
    payload = base64.b64decode(handle.read())
with open(path, "wb") as handle:
    handle.write(payload)

report = {"kind": kind, "opened": False}

if kind == "pptx":
    from pptx import Presentation

    presentation = Presentation(path)
    report["opened"] = True
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks = []
        overflow = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if text:
                blocks.append(text)
            width = getattr(shape, "width", None)
            height = getattr(shape, "height", None)
            if width and height:
                square_inches = (width / 914400.0) * (height / 914400.0)
                capacity = int(square_inches * 130)
                if capacity > 0 and len(text) > capacity:
                    overflow.append({"chars": len(text), "capacity": capacity})
        slides.append(
            {
                "index": index,
                "text_blocks": len(blocks),
                "chars": sum(len(block) for block in blocks),
                "overflow": overflow,
            }
        )
    report["slide_count"] = len(slides)
    report["slides"] = slides

elif kind == "docx":
    import docx

    document = docx.Document(path)
    report["opened"] = True
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    report["paragraph_count"] = len(paragraphs)
    report["heading_count"] = len(
        [p for p in paragraphs if (p.style is not None and "Heading" in str(p.style.name))]
    )
    report["chars"] = sum(len(p.text) for p in paragraphs)
    report["table_count"] = len(document.tables)

elif kind == "xlsx":
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    report["opened"] = True
    sheets = []
    formulas = 0
    for worksheet in workbook.worksheets:
        rows = 0
        cells = 0
        for row in worksheet.iter_rows():
            populated = [cell for cell in row if cell.value not in (None, "")]
            if populated:
                rows += 1
                cells += len(populated)
            for cell in populated:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
        sheets.append({"name": worksheet.title, "rows": rows, "cells": cells})
    report["sheet_count"] = len(sheets)
    report["sheets"] = sheets
    report["formula_count"] = formulas

print(json.dumps(report))
'''


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class ArtifactChecker:
    name = "artifacts"

    async def check(self, payload: VerificationInput) -> CheckResult:
        started = time.perf_counter()
        artifacts = list(payload.artifacts or [])
        if not artifacts:
            return CheckResult(
                checker=self.name,
                status=CheckStatus.SKIPPED,
                message="the result contains no artifacts to inspect",
                duration_ms=(time.perf_counter() - started) * 1000,
            )

        findings: list[dict[str, Any]] = []
        worst = CheckStatus.PASSED

        def escalate(status: CheckStatus) -> None:
            nonlocal worst
            order = {
                CheckStatus.PASSED: 0,
                CheckStatus.SKIPPED: 0,
                CheckStatus.WARNING: 1,
                CheckStatus.FAILED: 2,
            }
            if order[status] > order[worst]:
                worst = status

        for artifact in artifacts:
            artifact_id = artifact.get("artifact_id") or artifact.get("id")
            kind = str(artifact.get("type") or "").lower()
            storage_path = artifact.get("storage_path")
            base = {"artifact_id": artifact_id, "type": kind}

            if not storage_path:
                escalate(CheckStatus.FAILED)
                findings.append({**base, "type_of_finding": "no_storage_path"})
                continue

            path = Path(str(storage_path))
            if not path.is_file():
                escalate(CheckStatus.FAILED)
                findings.append({**base, "type_of_finding": "artifact_missing"})
                continue

            size = path.stat().st_size
            if size == 0:
                escalate(CheckStatus.FAILED)
                findings.append({**base, "type_of_finding": "artifact_empty"})
                continue

            recorded_size = artifact.get("size_bytes")
            if isinstance(recorded_size, int) and recorded_size != size:
                escalate(CheckStatus.FAILED)
                findings.append(
                    {
                        **base,
                        "type_of_finding": "size_mismatch",
                        "recorded": recorded_size,
                        "on_disk": size,
                    }
                )

            recorded_digest = str(artifact.get("sha256") or "")
            if recorded_digest:
                actual = _sha256(path)
                if actual != recorded_digest:
                    escalate(CheckStatus.FAILED)
                    findings.append(
                        {
                            **base,
                            "type_of_finding": "integrity_mismatch",
                            "recorded_sha256": recorded_digest[:16],
                            "actual_sha256": actual[:16],
                        }
                    )
            else:
                escalate(CheckStatus.WARNING)
                findings.append({**base, "type_of_finding": "no_recorded_digest"})

            structural = self._structural(path, kind)
            for item in structural.get("findings", []):
                escalate(CheckStatus(item.pop("status")))
                findings.append({**base, **item})

            if kind not in _OOXML_TYPES:
                continue

            deep = await self._deep_inspect(
                path=path,
                kind=kind,
                sandbox=payload.sandbox,
                job_id=payload.job_id,
                size=size,
            )
            for item in deep:
                escalate(CheckStatus(item.pop("status")))
                findings.append({**base, **item})

            expected = artifact.get("spec") or {}
            requested = expected.get("requested_units") if isinstance(expected, dict) else None
            produced = structural.get("units")
            if isinstance(requested, int) and isinstance(produced, int) and produced < requested:
                escalate(CheckStatus.FAILED)
                findings.append(
                    {
                        **base,
                        "type_of_finding": "fewer_units_than_requested",
                        "requested": requested,
                        "produced": produced,
                    }
                )

        message = {
            CheckStatus.PASSED: f"{len(artifacts)} artifact(s) opened and matched their recorded digest",
            CheckStatus.WARNING: f"{len(artifacts)} artifact(s) exist but could not be fully inspected",
            CheckStatus.FAILED: "at least one artifact is missing, altered, malformed or empty",
            CheckStatus.SKIPPED: "no artifact inspection was possible",
        }[worst]

        return CheckResult(
            checker=self.name,
            status=worst,
            message=message,
            findings=findings[:60],
            duration_ms=(time.perf_counter() - started) * 1000,
        )

    # --- host layer -------------------------------------------------------

    def _structural(self, path: Path, kind: str) -> dict[str, Any]:
        """Container-level checks using only the standard library."""
        findings: list[dict[str, Any]] = []
        units: int | None = None

        if kind in _OOXML_TYPES:
            if not zipfile.is_zipfile(path):
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "not_an_ooxml_container"}
                )
                return {"findings": findings, "units": units}
            try:
                with zipfile.ZipFile(path) as archive:
                    broken = archive.testzip()
                    if broken is not None:
                        findings.append(
                            {
                                "status": CheckStatus.FAILED.value,
                                "type_of_finding": "corrupt_container",
                                "member": broken,
                            }
                        )
                    names = archive.namelist()
                    required = _OOXML_PARTS[kind]
                    if required not in names:
                        findings.append(
                            {
                                "status": CheckStatus.FAILED.value,
                                "type_of_finding": "missing_ooxml_part",
                                "expected": required,
                            }
                        )
                    if kind == "pptx":
                        units = len(
                            [
                                name
                                for name in names
                                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
                            ]
                        )
                    elif kind == "xlsx":
                        units = len(
                            [
                                name
                                for name in names
                                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                            ]
                        )
                    if kind in {"pptx", "xlsx"} and units == 0:
                        findings.append(
                            {
                                "status": CheckStatus.FAILED.value,
                                "type_of_finding": "no_content_units",
                            }
                        )
            except zipfile.BadZipFile:
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "corrupt_container"}
                )

        elif kind == "pdf":
            head = path.open("rb").read(1024)
            data = path.read_bytes()
            if not head.startswith(b"%PDF-"):
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "not_a_pdf"}
                )
            if b"%%EOF" not in data[-2048:]:
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "truncated_pdf"}
                )
            units = len(_PDF_PAGE.findall(data))
            if units == 0:
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "pdf_has_no_pages"}
                )
            findings.append(
                {
                    "status": CheckStatus.WARNING.value,
                    "type_of_finding": "pdf_deep_inspection_unavailable",
                    "pages_structural": units,
                    "hint": "the documents image carries a PDF writer, not a reader",
                }
            )

        return {"findings": findings, "units": units}

    # --- sandbox layer ----------------------------------------------------

    async def _deep_inspect(
        self,
        *,
        path: Path,
        kind: str,
        sandbox: Any,
        job_id: str | None,
        size: int,
    ) -> list[dict[str, Any]]:
        if sandbox is None:
            return [
                {
                    "status": CheckStatus.WARNING.value,
                    "type_of_finding": "deep_inspection_unavailable",
                    "hint": "no sandbox is configured, so the file was never opened by a library",
                }
            ]
        if size > _MAX_SANDBOX_INSPECT_BYTES:
            return [
                {
                    "status": CheckStatus.WARNING.value,
                    "type_of_finding": "deep_inspection_skipped_too_large",
                    "size_bytes": size,
                    "limit_bytes": _MAX_SANDBOX_INSPECT_BYTES,
                }
            ]

        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        script = f"import sys\nsys.argv = ['inspect', {kind!r}]\n" + _INSPECT_SCRIPT
        try:
            execution = await sandbox.run_python(
                script,
                job_id=job_id,
                inputs={"artifact.b64": encoded},
            )
        except Exception as exc:
            return [
                {
                    "status": CheckStatus.WARNING.value,
                    "type_of_finding": "deep_inspection_error",
                    "error": f"{type(exc).__name__}: {exc}"[:200],
                }
            ]

        if not execution.get("ok"):
            return [
                {
                    "status": CheckStatus.FAILED.value,
                    "type_of_finding": "artifact_does_not_open",
                    "exit_code": execution.get("exit_code"),
                    "stderr": (execution.get("stderr") or "")[-400:],
                }
            ]

        try:
            report = json.loads((execution.get("stdout") or "").strip().splitlines()[-1])
        except (ValueError, IndexError):
            return [
                {
                    "status": CheckStatus.WARNING.value,
                    "type_of_finding": "deep_inspection_unreadable",
                    "hint": "the inspection script produced no parseable report",
                }
            ]

        return self._interpret(report, kind)

    def _interpret(self, report: dict[str, Any], kind: str) -> list[dict[str, Any]]:
        findings: list[dict[str, Any]] = []
        if not report.get("opened"):
            return [
                {
                    "status": CheckStatus.FAILED.value,
                    "type_of_finding": "artifact_does_not_open",
                }
            ]

        if kind == "pptx":
            slides = report.get("slides") or []
            if not slides:
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "deck_has_no_slides"}
                )
            empty = [slide["index"] for slide in slides if not slide.get("text_blocks")]
            if empty:
                findings.append(
                    {
                        "status": CheckStatus.FAILED.value,
                        "type_of_finding": "empty_slides",
                        "slides": empty[:20],
                    }
                )
            overflowing = [
                slide["index"] for slide in slides if slide.get("overflow")
            ]
            if overflowing:
                findings.append(
                    {
                        "status": CheckStatus.WARNING.value,
                        "type_of_finding": "possible_text_overflow",
                        "slides": overflowing[:20],
                        "hint": "estimated from shape area; PowerPoint may still shrink the text",
                    }
                )

        elif kind == "docx":
            if not report.get("paragraph_count"):
                findings.append(
                    {
                        "status": CheckStatus.FAILED.value,
                        "type_of_finding": "document_has_no_text",
                    }
                )

        elif kind == "xlsx":
            sheets = report.get("sheets") or []
            if not sheets:
                findings.append(
                    {"status": CheckStatus.FAILED.value, "type_of_finding": "workbook_has_no_sheets"}
                )
            empty = [sheet["name"] for sheet in sheets if not sheet.get("rows")]
            if empty:
                findings.append(
                    {
                        "status": CheckStatus.FAILED.value,
                        "type_of_finding": "empty_sheets",
                        "sheets": empty[:20],
                    }
                )
            if report.get("formula_count"):
                findings.append(
                    {
                        "status": CheckStatus.WARNING.value,
                        "type_of_finding": "formulas_present",
                        "count": report["formula_count"],
                        "hint": "the generator writes values, not formulas; investigate the source",
                    }
                )

        return findings
