"""Word generator - runs *inside* the sandbox, never on the host.

Invoked as::

    python docx_generator.py --spec inputs/spec.json --out artifacts/report.docx

Reads a validated ``backend.deliverables.spec.ArtifactSpec`` as JSON and emits
a .docx. Imports nothing from ``backend`` - the container holds only
python-docx and the standard library.

Same rule as the other generators: the model supplies the *spec*, this
reviewed script supplies the *code*. Spec values reach python-docx only as
strings, so a heading containing braces, quotes or ``{{`` is content rather
than a format string.

A missing python-docx exits 3 with an actionable message instead of
traceback-ing. The sandbox is network-denied by policy, so ``pip install``
cannot be the fix - the documents image has to be built with the library
baked in (see infrastructure/docker/sandbox-documents.Dockerfile).
"""

from __future__ import annotations

import argparse
import json
import sys

#: Word tolerates long paragraphs, but a runaway one is a sign the spec was
#: built from unbounded model output. Clip and keep going.
MAX_PARAGRAPH_CHARS = 4000


def _citation_suffix(citations: list | None) -> str:
    """Render citations as ``[doc p.12; doc p.13]``.

    Inline rather than as footnotes: python-docx cannot create real footnotes
    without hand-writing OOXML, and a fake footnote that looks real is worse
    than a visible bracket.
    """
    labels = []
    for citation in citations or []:
        document_id = str(citation.get("document_id", "") or "")
        if not document_id:
            continue
        page = citation.get("page")
        section = str(citation.get("section", "") or "")
        label = f"{document_id} p.{page}" if page else document_id
        if section:
            label += f" ({section})"
        labels.append(label)
    if not labels:
        return ""
    return "  [" + "; ".join(labels[:3]) + "]"


def _add_section(document, section: dict) -> int:
    """Write one section. Returns the number of paragraphs written."""
    written = 0
    level = int(section.get("level", 1) or 1)
    level = 1 if level < 1 else 4 if level > 4 else level
    document.add_heading(str(section.get("heading", ""))[:300], level=level)

    for paragraph in section.get("paragraphs") or []:
        text = str(paragraph)[:MAX_PARAGRAPH_CHARS]
        if not text.strip():
            continue
        document.add_paragraph(text)
        written += 1

    for bullet in section.get("bullets") or []:
        text = str(bullet.get("text", ""))[:MAX_PARAGRAPH_CHARS]
        if not text.strip():
            continue
        document.add_paragraph(text + _citation_suffix(bullet.get("citations")), style="List Bullet")
        written += 1

    trailing = _citation_suffix(section.get("citations"))
    if trailing:
        document.add_paragraph(f"Sources for this section:{trailing}")
        written += 1
    return written


def _add_sources(document, sources: list) -> int:
    """Numbered source list. Every generator emits one.

    An artifact that cannot be traced back to pages is not a deliverable in
    this product, so the list is written even when the model forgot to ask
    for it.
    """
    if not sources:
        return 0
    document.add_heading("Sources", level=1)
    count = 0
    for citation in sources:
        document_id = str(citation.get("document_id", "") or "")
        if not document_id:
            continue
        page = citation.get("page")
        section = str(citation.get("section", "") or "")
        line = f"{document_id}"
        if page:
            line += f", page {page}"
        if section:
            line += f" - {section}"
        document.add_paragraph(line, style="List Number")
        count += 1
    return count


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Render an artifact spec as .docx")
    parser.add_argument("--spec", required=True, help="path to the JSON artifact spec")
    parser.add_argument("--out", required=True, help="path to write the .docx to")
    args = parser.parse_args(argv)

    try:
        from docx import Document
    except ImportError:
        print(
            "python-docx is not installed in this sandbox image. The sandbox has no "
            "network access by policy, so it cannot be installed at runtime: rebuild "
            "the documents image (infrastructure/docker/sandbox-documents.Dockerfile) "
            "with python-docx baked in.",
            file=sys.stderr,
        )
        return 3

    with open(args.spec, encoding="utf-8") as handle:
        spec = json.load(handle)

    document = Document()
    document.core_properties.title = str(spec.get("title", "Untitled"))[:300]
    document.core_properties.author = str(spec.get("author", "Project 117"))[:120]

    document.add_heading(str(spec.get("title", "Untitled"))[:300], level=0)
    subtitle = str(spec.get("subtitle", "") or "")
    if subtitle:
        document.add_paragraph(subtitle)

    sections = spec.get("sections") or []
    paragraphs = 0
    for section in sections:
        paragraphs += _add_section(document, section)

    sources = _add_sources(document, spec.get("sources") or [])

    footer = str(spec.get("footer", "") or "")
    if footer:
        document.add_paragraph(footer)

    document.save(args.out)

    # stdout is parsed by the calling tool; keep it machine-readable and free
    # of document content.
    print(
        json.dumps(
            {
                "artifact": args.out,
                "sections": len(sections),
                "paragraphs": paragraphs,
                "sources": sources,
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
