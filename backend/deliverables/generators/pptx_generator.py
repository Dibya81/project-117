"""PowerPoint generator - runs *inside* the sandbox, never on the host.

Uploaded into the sandbox as a file and invoked as::

    python pptx_generator.py --spec inputs/spec.json --out artifacts/deck.pptx

It reads a validated ``backend.deliverables.spec.ArtifactSpec`` as JSON and
emits a .pptx. It imports nothing from ``backend``: it has to run in a
container that contains only python-pptx and the standard library.

Why a static script instead of model-generated python-pptx code: this file was
reviewed once and behaves identically on every run. Generated code would have
to be reviewed on every run, which nobody does. The model's creativity belongs
in the *content* of the spec, not in the code path that renders it.

Spec values are only ever passed to python-pptx as strings - no ``eval``, no
format strings built from spec values, no shelling out - so a slide title
containing braces or quotes is content, not injection.
"""

from __future__ import annotations

import argparse
import json
import sys

#: Rough capacity of one content slide before text starts overflowing. The
#: artifact checker re-derives this independently; here we only avoid
#: producing something obviously broken.
BULLETS_PER_SLIDE = 8
CHARS_PER_BULLET = 240


def _citation_suffix(citations: list | None) -> str:
    labels = []
    for citation in citations or []:
        document_id = str(citation.get("document_id", ""))
        page = citation.get("page")
        if not document_id:
            continue
        labels.append(f"{document_id} p.{page}" if page else document_id)
    if not labels:
        return ""
    return "  [" + "; ".join(labels[:3]) + "]"


def _add_title_slide(prs, spec: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = str(spec.get("title", "Untitled"))
    if len(slide.placeholders) > 1:
        subtitle = str(spec.get("subtitle", "") or "")
        slide.placeholders[1].text = subtitle or str(spec.get("author", "") or "")


def _add_content_slide(prs, slide_spec: dict) -> None:
    bullets = slide_spec.get("bullets") or []
    layout_name = str(slide_spec.get("layout", "title_content"))
    # Layout 5 is "Title Only"; layout 1 is "Title and Content".
    layout_index = 5 if layout_name == "section" or not bullets else 1
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    slide.shapes.title.text = str(slide_spec.get("title", ""))

    if bullets and len(slide.placeholders) > 1:
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(bullets[:BULLETS_PER_SLIDE]):
            text = str(bullet.get("text", ""))[:CHARS_PER_BULLET]
            text += _citation_suffix(bullet.get("citations"))
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.level = 0

    notes = str(slide_spec.get("notes", "") or "")
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_sources_slide(prs, sources: list) -> None:
    if not sources:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sources"
    frame = slide.placeholders[1].text_frame
    frame.clear()
    for index, citation in enumerate(sources[:20]):
        document_id = str(citation.get("document_id", ""))
        page = citation.get("page")
        section = str(citation.get("section", "") or "")
        label = f"{document_id} - p. {page}" if page else document_id
        if section:
            label += f" - {section}"
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = label


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Render an ArtifactSpec as .pptx")
    parser.add_argument("--spec", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args(argv)

    try:
        from pptx import Presentation
    except ImportError:
        print(
            "python-pptx is not installed in this sandbox image. Build the documents "
            "image (infrastructure/docker/sandbox-documents.Dockerfile) and set "
            "P117_SANDBOX_DOCUMENTS_IMAGE.",
            file=sys.stderr,
        )
        return 3

    with open(args.spec, encoding="utf-8") as handle:
        spec = json.load(handle)

    prs = Presentation()
    _add_title_slide(prs, spec)
    for slide_spec in spec.get("slides") or []:
        _add_content_slide(prs, slide_spec)
    _add_sources_slide(prs, spec.get("sources") or [])
    prs.save(args.out)

    # Machine-readable last line: the tool parses this instead of scraping
    # arbitrary stdout.
    print(json.dumps({"artifact": args.out, "slides": len(prs.slides._sldIdLst)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
